"""Extract text and structure from Word (.docx) files."""

import logging
from pathlib import Path

log = logging.getLogger(__name__)


class DocxParser:
    """Parse .docx files into Markdown-ready text using python-docx."""

    def parse(self, filepath: Path) -> str:
        from docx import Document

        log.info("Parsing DOCX: %s", filepath)
        doc = Document(str(filepath))
        lines: list[str] = []

        for para in doc.paragraphs:
            text = para.text.strip()
            if not text:
                lines.append("")
                continue

            style = (para.style.name or "").lower()

            # Map heading styles to Markdown headings
            if style.startswith("heading"):
                try:
                    level = int(style.replace("heading", "").strip())
                except ValueError:
                    level = 1
                level = min(level, 6)
                lines.append(f"{'#' * level} {text}")
            elif style.startswith("list"):
                lines.append(f"- {text}")
            else:
                lines.append(text)

        # Extract tables
        for table_idx, table in enumerate(doc.tables):
            lines.append("")
            lines.append(f"**Table {table_idx + 1}:**")
            lines.append("")
            for row_idx, row in enumerate(table.rows):
                cells = [cell.text.strip().replace("|", "\\|") for cell in row.cells]
                lines.append("| " + " | ".join(cells) + " |")
                if row_idx == 0:
                    lines.append("| " + " | ".join(["---"] * len(cells)) + " |")

        return "\n".join(lines)
